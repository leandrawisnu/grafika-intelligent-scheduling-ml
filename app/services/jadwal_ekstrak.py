import base64
import json
from io import BytesIO

import pymupdf
from langchain_core.messages import HumanMessage
from langchain_openai import ChatOpenAI
from openpyxl import load_workbook

from app.config import LLM_MODEL, OPENAI_API_KEY
from app.models.schemas import BarisJadwalEkstrak

AMBANG_TEKS_PDF = 80
MAKS_HALAMAN = 6
MAKS_BARIS_EXCEL = 250
MAKS_KARAKTER = 30000

PROMPT = """Ekstrak jadwal mengajar menjadi JSON.
Balas hanya JSON dengan bentuk {"baris":[{"hari":"","jam":"","mata_pelajaran":"","kelas":"","guru":"","ruangan":""}]}.
Satu objek sama dengan satu jam pelajaran.
Lewati istirahat, upacara, dan sel kosong.
Jam berisi angka jam ke bila terlihat, atau rentang waktu.
Gunakan ejaan dari katalog bila namanya mirip.
Jangan menulis UUID.
"""


def excel_ke_teks(data: bytes) -> str:
    buku = load_workbook(BytesIO(data), read_only=True, data_only=True)
    bagian: list[str] = []
    try:
        for lembar in buku.worksheets:
            baris: list[str] = []
            for row in lembar.iter_rows(values_only=True):
                sel = ["" if nilai is None else str(nilai).strip() for nilai in row]
                if any(sel):
                    baris.append(" | ".join(sel))
                if len(baris) >= MAKS_BARIS_EXCEL:
                    break
            if baris:
                bagian.append(f"# {lembar.title}\n" + "\n".join(baris))
    finally:
        buku.close()
    teks = "\n\n".join(bagian).strip()
    if not teks:
        raise ValueError("Excel tidak berisi tabel")
    return teks[:MAKS_KARAKTER]


def pdf_ke_muatan(data: bytes) -> tuple[str, list[bytes]]:
    dokumen = pymupdf.open(stream=data, filetype="pdf")
    try:
        jumlah = min(dokumen.page_count, MAKS_HALAMAN)
        if jumlah == 0:
            raise ValueError("PDF tidak berisi halaman")
        potongan = [dokumen[i].get_text("text") or "" for i in range(jumlah)]
        teks = "\n".join(potongan).strip()
        if len(teks) >= AMBANG_TEKS_PDF:
            return teks[:MAKS_KARAKTER], []
        gambar: list[bytes] = []
        for i in range(jumlah):
            halaman = dokumen[i]
            matriks = pymupdf.Matrix(1.3, 1.3)
            pix = halaman.get_pixmap(matrix=matriks, alpha=False)
            if pix.width > 1600:
                skala = 1600 / pix.width
                pix = halaman.get_pixmap(matrix=pymupdf.Matrix(1.3 * skala, 1.3 * skala), alpha=False)
            gambar.append(pix.tobytes("png"))
        return teks[:MAKS_KARAKTER], gambar
    finally:
        dokumen.close()


def parse_baris(konten: str) -> list[BarisJadwalEkstrak]:
    teks = konten.strip()
    if teks.startswith("```"):
        teks = teks.split("\n", 1)[1].rsplit("```", 1)[0].strip()
    data = json.loads(teks)
    mentah = data if isinstance(data, list) else data.get("baris", [])
    if not isinstance(mentah, list):
        raise ValueError("Jawaban model bukan daftar baris")
    hasil: list[BarisJadwalEkstrak] = []
    for item in mentah:
        if not isinstance(item, dict):
            continue
        hasil.append(
            BarisJadwalEkstrak(
                hari=str(item.get("hari") or ""),
                jam=str(item.get("jam") or ""),
                mata_pelajaran=str(item.get("mata_pelajaran") or ""),
                kelas=str(item.get("kelas") or ""),
                guru=str(item.get("guru") or ""),
                ruangan=str(item.get("ruangan") or ""),
            )
        )
    return hasil


def ringkas_katalog(katalog: dict) -> str:
    baris: list[str] = []
    for kunci in ("hari", "jam", "mata_pelajaran", "kelas", "guru", "ruangan"):
        nilai = katalog.get(kunci) or []
        if not isinstance(nilai, list):
            continue
        nama = [str(item) for item in nilai[:300] if str(item).strip()]
        if nama:
            baris.append(f"{kunci}: {', '.join(nama)}")
    return "\n".join(baris)


def susun_baris(teks: str, katalog: dict, gambar: list[bytes] | None) -> list[BarisJadwalEkstrak]:
    if not OPENAI_API_KEY or OPENAI_API_KEY == "sk-placeholder":
        raise ValueError("OPENAI_API_KEY belum diisi")
    llm = ChatOpenAI(api_key=OPENAI_API_KEY, model=LLM_MODEL, temperature=0)
    isi: list[dict] = [
        {"type": "text", "text": PROMPT + "\nKatalog:\n" + ringkas_katalog(katalog) + "\n\nIsi berkas:\n" + teks}
    ]
    for png in gambar or []:
        kode = base64.b64encode(png).decode("ascii")
        isi.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{kode}"}})
    jawab = llm.invoke([HumanMessage(content=isi)])
    konten = jawab.content if isinstance(jawab.content, str) else str(jawab.content)
    return parse_baris(konten)
